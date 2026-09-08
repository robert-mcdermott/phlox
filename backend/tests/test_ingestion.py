"""Wave 7 real parsing, durable queue, publication, provider-failure and privacy contracts."""
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from sqlalchemy import inspect

from app.models import DocChunk, Document, Setting
from app.rag import identity, jobs, maintenance, retrieve
from app.rag.parsing import chunks
from app.rag.store import QdrantVectorStore


@pytest.fixture
def library(db, tmp_path, monkeypatch):
    store = QdrantVectorStore({'path': str(tmp_path / 'qdrant')})
    previous = maintenance.state(db)
    # Other suites intentionally leave legacy documents in the shared test DB. A library
    # migration is global, so isolate this fixture's library from those unrelated rows.
    previous_docs = [(d.id, d.status) for d in db.query(Document).all()]
    for key, _ in previous_docs:
        db.get(Document, key).status = 'fixture-hidden'
    db.commit()
    maintenance.save(db, {})
    for module in ('app.rag.store', 'app.rag.retrieve', 'app.rag.maintenance'):
        monkeypatch.setattr(module + '.get_vector_store', lambda: store)
    worker = jobs.Worker()
    monkeypatch.setattr(jobs, 'worker', worker)
    docs = []
    yield SimpleNamespace(store=store, worker=worker, docs=docs)
    db.rollback()
    for row in docs:
        key = row if isinstance(row, str) else inspect(row).identity[0]
        current = db.get(Document, key, populate_existing=True)
        if current:
            db.delete(current)
    db.commit()
    for key, status in previous_docs:
        row = db.get(Document, key)
        if row:
            row.status = status
    db.commit()
    maintenance.save(db, previous)
    store._client.close()


def upload(client, db, library, text=b'Vehicle policy: bicycles are allowed.', filename='policy.txt', process=True):
    response = client.post('/api/documents', files={'file': (filename, text, 'application/octet-stream')})
    assert response.status_code == 200, response.text
    data = response.json()
    library.docs.append(data['id'])
    assert data['status'] == 'queued'
    if process:
        assert library.worker.step()
    return db.get(Document, data['id'], populate_existing=True)


def test_queue_progress_retry_and_restart_do_not_duplicate_chunks(client, db, library):
    doc = upload(client, db, library, process=False)
    token = doc.ingestion['token']
    assert client.post(f'/api/documents/{doc.id}/retry').status_code == 409
    library.worker.recover()
    db.refresh(doc)
    assert doc.status == 'interrupted'
    assert client.post(f'/api/documents/{doc.id}/retry').status_code == 200
    db.refresh(doc)
    assert doc.ingestion['token'] != token
    assert library.worker.step()
    db.refresh(doc)
    assert doc.status == 'ready' and doc.ingestion['completed'] == doc.n_chunks == 1
    assert client.post(f'/api/documents/{doc.id}/retry').status_code == 200
    assert library.worker.step()
    db.refresh(doc)
    assert db.query(DocChunk).filter_by(document_id=doc.id).count() == doc.n_chunks == 1
    assert doc.ingestion['attempt'] == 3
    found = retrieve.search_chunks(db, 'bicycles', user_id='local')
    assert found[0]['document_id'] == doc.id


def test_provider_failure_never_inserts_hash_vectors_and_can_retry(client, db, library, monkeypatch):
    cfg = {'profile': 'embedding', 'model': 'semantic', 'version': '1'}
    monkeypatch.setattr(identity, 'get_embeddings_config', lambda: cfg)
    monkeypatch.setattr(identity, 'get_profile', lambda _: {'type': 'openai', 'endpoint': 'http://embedding.invalid/v1'})
    class Provider:
        fail = True
        def embed(self, texts, model=None):
            if self.fail:
                raise RuntimeError('secret-provider-token must not leak')
            return [[1.0, 0.0] for _ in texts]
    provider = Provider()
    monkeypatch.setattr('app.providers.registry.build_provider', lambda _: provider)
    doc = upload(client, db, library)
    assert doc.status == 'error' and 'secret' not in doc.error
    assert db.query(DocChunk).filter_by(document_id=doc.id).count() == 0
    provider.fail = False
    client.post(f'/api/documents/{doc.id}/retry')
    library.worker.step()
    db.refresh(doc)
    assert doc.status == 'ready' and doc.chunks[0].embedding_identity['dimensions'] == 2
    provider.fail = True
    results = retrieve.search_chunks(db, 'bicycles', user_id='local')
    assert results and 'Keyword search only' in results.notice
    assert doc.chunks[0].embedding == [1.0, 0.0]
    assert not retrieve.search_chunks(db, 'unmatchedword', user_id='local')


def test_same_dimension_change_requires_rebuild_and_failed_stage_preserves_index(client, db, library, monkeypatch):
    cfg = {'profile': 'embedding', 'model': 'model-a'}
    monkeypatch.setattr(identity, 'get_embeddings_config', lambda: cfg)
    monkeypatch.setattr(identity, 'get_profile', lambda _: {'type': 'openai', 'endpoint': 'http://embedding.invalid/v1'})
    monkeypatch.setattr('app.providers.registry.build_provider', lambda _: SimpleNamespace(embed=lambda texts, model=None: [[1.0, 0.0] for t in texts]))
    doc = upload(client, db, library)
    first_spec = doc.chunks[0].embedding_identity.copy()
    old_collection = library.store.collection
    cfg['model'] = 'model-b'
    assert maintenance.status(db)['rebuild_required']
    assert retrieve.search_chunks(db, 'bicycles', user_id='local').notice
    original = library.store.stage
    monkeypatch.setattr(library.store, 'stage', lambda *a, **k: (_ for _ in ()).throw(RuntimeError('index offline')))
    assert client.post('/api/documents/reindex').status_code == 200
    library.worker.step()
    db.expire_all()
    assert maintenance.state(db)['status'] == 'error'
    assert doc.chunks[0].embedding_identity == first_spec
    assert library.store.collection == old_collection
    monkeypatch.setattr(library.store, 'stage', original)
    client.post('/api/documents/reindex')
    library.worker.step()
    db.expire_all()
    assert maintenance.state(db)['status'] == 'ready'
    assert not maintenance.status(db)['rebuild_required']
    assert doc.chunks[0].embedding_identity['model'] == 'model-b'
    assert library.store.collection != old_collection
    # The SQL pointer restores the successfully published collection after process restart.
    assert db.get(Setting, 'rag:index').value['collection'] == library.store.collection


def test_incomplete_embeddings_and_deletion_during_processing(client, db, library, monkeypatch):
    original = identity.embed
    monkeypatch.setattr('app.rag.ingest.embed', lambda *a, **k: ([], {'dimensions': 2}))
    doc = upload(client, db, library)
    assert doc.status == 'error' and not doc.chunks
    doc_id = doc.id
    def delete_while_embedding(texts, expected):
        assert client.delete(f'/api/documents/{doc_id}').status_code == 200
        return original(texts, expected)
    monkeypatch.setattr('app.rag.ingest.embed', delete_while_embedding)
    client.post(f'/api/documents/{doc.id}/retry')
    library.worker.step()
    assert db.get(Document, doc_id, populate_existing=True) is None
    assert db.query(DocChunk).filter_by(document_id=doc_id).count() == 0


def pdf_fixture(path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for text in ['First page: bicycles permitted.', 'Second page: helmets required.']:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 10 200 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = stream
    writer.write(path)


def test_pdf_pages_docx_tables_and_markdown_sections_are_real_locations(tmp_path):
    path = tmp_path / 'pages.pdf'
    pdf_fixture(path)
    passages, _ = chunks(path)
    assert [p[1]['page'] for p in passages] == [1, 2]
    assert 'helmets' in passages[1][0]
    from docx import Document as WordDocument
    doc = WordDocument()
    doc.add_heading('Travel policy', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Meal limit', '$45'
    doc.save(tmp_path / 'policy.docx')
    passages, _ = chunks(tmp_path / 'policy.docx')
    assert any('Meal limit | $45' in text and loc['table_row'] == 1 and loc['section'] == 'Travel policy' for text, loc in passages)
    (tmp_path / 'policy.md').write_text('# Travel\n\nMeal allowance is $45.')
    assert chunks(tmp_path / 'policy.md')[0][0][1]['section'] == 'Travel'


def test_pdf_citations_keep_page_in_panel_and_export(client, db, library, tmp_path):
    from app.models import Conversation, Message
    from app.sources import capture, export_markdown, inspect_source
    path = tmp_path / 'pages.pdf'
    pdf_fixture(path)
    doc = upload(client, db, library, path.read_bytes(), 'pages.pdf')
    conv = Conversation(user_id='local', title='Page evidence')
    db.add(conv)
    db.commit()
    try:
        chunk = db.query(DocChunk).filter_by(document_id=doc.id, ordinal=1).one()
        ref, text = capture(db, conversation_id=conv.id, user_id='local', turn_id='page-turn', document_id=doc.id, chunk_id=chunk.id)
        assert 'page 2' in text
        db.add(Message(conversation_id=conv.id, role='assistant', content='Wear helmets [S1].', citations=[ref]))
        db.commit()
        assert inspect_source(db, conv, ref['source_id'])['location']['page'] == 2
        assert 'page 2' in export_markdown(db, conv)
    finally:
        db.delete(conv)
        db.commit()


def test_unsupported_binary_and_limits_are_visible(client, db, library, monkeypatch):
    doc = upload(client, db, library, b'\x00binary', 'file.exe')
    assert doc.status == 'error' and 'Unsupported' in doc.error
    monkeypatch.setattr('app.routers.documents.MAX_BYTES', 10)
    count = db.query(Document).count()
    assert client.post('/api/documents', files={'file': ('big.txt', b'x'*11, 'text/plain')}).status_code == 413
    assert db.query(Document).count() == count


def test_keyword_benchmark_and_private_distractors(client, db, library):
    fixture = json.loads((Path(__file__).parent.parent / 'evals/retrieval_v1.json').read_text())
    ids = {}
    for item in fixture['documents']:
        doc = upload(client, db, library, item['text'].encode(), item['name'] + '.txt')
        ids[item['name']] = doc.id
        if item.get('private'):
            doc.user_id = 'someone-else'
        db.commit()
    for chunk in db.query(DocChunk).filter(DocChunk.document_id.in_(ids.values())):
        chunk.embedding_identity = None
    db.commit()
    for case in fixture['queries']:
        found = retrieve.search_chunks(db, case['query'], user_id='local', top_k=5)
        assert found.notice and 'Keyword' in found.notice
        assert {r['document_id'] for r in found} == {ids[name] for name in case['documents']}
    assert client.post(f'/api/documents/{ids["private"]}/retry').status_code == 404


def test_staging_write_failure_keeps_previous_collection_searchable(client, db, library, monkeypatch):
    doc = upload(client, db, library)
    old = library.store.collection
    upsert = library.store._client.upsert
    def fail_staging(*args, **kwargs):
        if kwargs.get('collection_name') != old:
            raise RuntimeError('staging write failed')
        return upsert(*args, **kwargs)
    monkeypatch.setattr(library.store._client, 'upsert', fail_staging)
    client.post('/api/documents/reindex')
    library.worker.step()
    assert library.store.collection == old
    assert retrieve.search_chunks(db, 'bicycles', user_id='local')[0]['document_id'] == doc.id
    assert {c.name for c in library.store._client.get_collections().collections} == {old}


@pytest.mark.parametrize('response', [[], [[float('nan')]], [[1.0], [1.0, 2.0]]])
def test_provider_batch_validation_rejects_invalid_vectors(monkeypatch, response):
    monkeypatch.setattr(identity, 'get_embeddings_config', lambda: {'profile': 'p', 'model': 'm'})
    monkeypatch.setattr(identity, 'get_profile', lambda _: {'type': 'openai'})
    monkeypatch.setattr('app.providers.registry.build_provider', lambda _: SimpleNamespace(embed=lambda *a, **k: response))
    with pytest.raises(identity.EmbeddingError):
        identity.embed(['first', 'second'])


def test_semantic_paraphrase_fixture_with_scripted_embedding_provider(client, db, library, monkeypatch):
    monkeypatch.setattr(identity, 'get_embeddings_config', lambda: {'profile': 'p', 'model': 'scripted'})
    monkeypatch.setattr(identity, 'get_profile', lambda _: {'type': 'openai'})
    def embedding(texts, model=None):
        return [[1.0, 0.0] if any(word in text.lower() for word in ['bicycles', 'cycling']) else [0.0, 1.0] for text in texts]
    monkeypatch.setattr('app.providers.registry.build_provider', lambda _: SimpleNamespace(embed=embedding))
    desired = upload(client, db, library, b'Bicycles require helmets.', 'transport.txt')
    upload(client, db, library, b'Meal reimbursement is forty dollars.', 'travel.txt')
    found = retrieve.search_chunks(db, 'cycling', user_id='local', top_k=1)
    assert found[0]['document_id'] == desired.id and found.notice is None


def test_docx_table_retrieval_improves_over_previous_paragraph_only_extraction(client, db, library, tmp_path):
    from docx import Document as WordDocument
    word = WordDocument()
    word.add_heading('Travel policy', level=1)
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Reimbursement', '$45'
    path = tmp_path / 'table.docx'
    word.save(path)
    legacy = '\n'.join(p.text for p in WordDocument(path).paragraphs)
    assert 'Reimbursement' not in legacy  # Measured former extraction misses the fixture fact.
    doc = upload(client, db, library, path.read_bytes(), 'table.docx')
    found = retrieve.search_chunks(db, 'reimbursement', user_id='local', top_k=1)
    assert found[0]['document_id'] == doc.id and '$45' in found[0]['text']


def test_assistant_retry_scope_and_admin_index_controls(client, db, library):
    from app.models import Assistant, User
    from app.main import app
    from app.auth.deps import get_current_user
    assistant = Assistant(name='Processing KB', created_by='local', visibility='public')
    db.add(assistant)
    db.commit()
    try:
        response = client.post(f'/api/assistants/{assistant.id}/documents', files={'file': ('kb.txt', b'Bicycles are allowed.', 'text/plain')})
        assert response.status_code == 200
        doc_id = response.json()['id']
        library.docs.append(doc_id)
        library.worker.recover()
        assert client.post(f'/api/documents/{doc_id}/retry').status_code == 404
        assert client.post(f'/api/assistants/{assistant.id}/documents/{doc_id}/retry').status_code == 200
        library.worker.step()
        assert db.get(Document, doc_id, populate_existing=True).status == 'ready'
        app.dependency_overrides[get_current_user] = lambda: User(id='foreign', username='foreign', role='user')
        assert client.get('/api/documents/index-status').status_code == 403
        assert client.post('/api/documents/reindex').status_code == 403
        assert client.post(f'/api/documents/{doc_id}/retry').status_code == 404
    finally:
        app.dependency_overrides.pop(get_current_user, None)
        db.delete(assistant)
        db.commit()


def test_shared_visibility_revoked_during_embedding_is_rechecked(client, db, library, monkeypatch):
    from app.models import Assistant
    assistant = Assistant(name='Revoked KB', created_by='other', visibility='public')
    db.add(assistant)
    db.commit()
    try:
        doc = upload(client, db, library)
        doc.user_id, doc.assistant_id = None, assistant.id
        db.commit()
        original = identity.embed
        def revoke(*args, **kwargs):
            vectors = original(*args, **kwargs)
            assistant.visibility = 'private'
            db.commit()
            return vectors
        monkeypatch.setattr(identity, 'embed', revoke)
        assert not retrieve.search_chunks(db, 'bicycles', user_id='local', assistant_id=assistant.id)
    finally:
        db.delete(assistant)
        db.commit()


def test_rebuild_recovery_preserves_published_pointer(db, library, monkeypatch):
    maintenance.save(db, {'status': 'processing', 'collection': 'previous-good-index'})
    library.worker.recover()
    recovered = maintenance.state(db)
    assert recovered['status'] == 'interrupted'
    assert recovered['collection'] == 'previous-good-index'
    def limit_failure(*args):
        raise identity.EmbeddingError('Rebuild exceeds the 10,000 chunk limit.')
    monkeypatch.setattr(jobs, 'sync_index', limit_failure)
    jobs.queue_rebuild(db)
    library.worker.step()
    assert '10,000 chunk limit' in maintenance.state(db)['error']
    assert maintenance.state(db)['collection'] == 'previous-good-index'
